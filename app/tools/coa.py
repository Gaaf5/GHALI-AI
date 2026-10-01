from __future__ import annotations
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "coa_template.xlsx"

def _num(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default

def _trace_ppm(breakdown):
    direct = breakdown.get("te_mix_trace_elements_ppm") or {}
    if direct:
        return {str(k).lower().replace("_", "-"): _num(v) for k, v in direct.items()}
    pct = breakdown.get("trace_elements_pct_of_product") or {}
    return {str(k).lower().replace("_", "-"): _num(v) * 10000.0 for k, v in pct.items()}

def _trace_value(traces, *names):
    for name in names:
        key = name.lower().replace("_", "-")
        if key in traces:
            return traces[key]
    return None

def build_certificate_of_analysis(payload):
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"COA template not found: {TEMPLATE}")
    wb = load_workbook(TEMPLATE)
    if "Sheet2" not in wb.sheetnames:
        raise ValueError("COA template is missing Sheet2")
    ws = wb["Sheet2"]
    for name in list(wb.sheetnames):
        if name != "Sheet2":
            del wb[name]
    ws.title = "COA"

    achieved = payload.get("achieved") or {}
    breakdown = payload.get("nutrient_breakdown") or {}
    forms = breakdown.get("nitrogen_forms_pct_of_product") or {}
    traces = _trace_ppm(breakdown)

    n = _num(achieved.get("N"))
    p = _num(achieved.get("P2O5"))
    k = _num(achieved.get("K2O"))
    product_name = f"{n:.2f}-{p:.2f}-{k:.2f}"

    ws["B4"] = f" Product Name: {product_name}"
    ws["C7"] = n
    ws["C8"] = _num(forms.get("nitrate_N"))
    ws["C9"] = _num(forms.get("ammoniacal_N"))
    ws["C10"] = _num(forms.get("urea_N"))
    ws["C11"] = p
    ws["C12"] = k

    ws["C14"] = _trace_value(traces, "Fe-EDTA", "Fe EDTA", "Fe")
    ws["C15"] = _trace_value(traces, "Zn-EDTA", "Zn EDTA", "Zn")
    ws["C16"] = _trace_value(traces, "Mn-EDTA", "Mn EDTA", "Mn")
    ws["C17"] = _trace_value(traces, "Cu-EDTA", "Cu EDTA", "Cu")
    ws["C18"] = _trace_value(traces, "B", "Boron")

    cl = _num(breakdown.get("chlorine_pct_of_product"))
    if cl:
        ws["C26"] = f"{cl:.4f} %"
    # Reverse Analysis does not measure pH, so never leave a formula pointing
    # at the removed Formula sheet. Keep this field blank rather than inventing it.
    ws["C31"] = None
    # Keep the template's specification text for heavy metals, sodium,
    # appearance and solubility. Reverse Analysis has no measured values
    # for those properties, so they must not be fabricated.
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
