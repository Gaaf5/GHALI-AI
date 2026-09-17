from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


def load_docx(path: str | Path) -> str:
    """Extract paragraph and table text from a DOCX document."""
    if Document is None:
        raise RuntimeError("DOCX support requires the 'python-docx' package")

    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(file_path)

    document = Document(str(file_path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()
